from docx import Document as DocxDocument
from langchain_core.documents import Document  
def load_docx(path):
    doc = DocxDocument(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    return [Document(page_content=text, metadata={"source": path})]
